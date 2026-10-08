from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment


EXCEL_FILE = Path("application_tracker.xlsx")

HEADERS = [
    "Date",
    "Time",
    "Job Title",
    "Company",
    "Match Score",
    "AI Decision",
    "Application Status",
    "Naukri Job URL",
    "External URL",
    "Current URL",
    "Reason",
    "Notes",
    "Unanswered Question",
]


def create_excel_file():
    if EXCEL_FILE.exists():
        return

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Applications"

    for column, header in enumerate(HEADERS, start=1):
        cell = sheet.cell(row=1, column=column, value=header)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:M1"

    widths = {
        "A": 14,
        "B": 12,
        "C": 35,
        "D": 30,
        "E": 14,
        "F": 14,
        "G": 28,
        "H": 60,
        "I": 60,
        "J": 60,
        "K": 40,
        "L": 50,
        "M": 60,
    }

    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    workbook.save(EXCEL_FILE)


def prepare_real_tracker():
    create_excel_file()

    workbook = load_workbook(EXCEL_FILE)
    sheet = workbook["Applications"]

    # Make sure the new column exists even if the old Excel file
    # was created with the previous 12-column version.
    for column, header in enumerate(HEADERS, start=1):
        sheet.cell(row=1, column=column, value=header)

    # Remove development TEST rows only.
    rows_to_delete = []

    for row in range(2, sheet.max_row + 1):
        status = sheet.cell(row=row, column=7).value

        if status == "TEST":
            rows_to_delete.append(row)

    for row in reversed(rows_to_delete):
        sheet.delete_rows(row)

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:M1"

    workbook.save(EXCEL_FILE)

    if rows_to_delete:
        print(f"Removed {len(rows_to_delete)} old TEST row(s).")


def add_hyperlink(cell, url):
    if not url:
        return

    cell.hyperlink = url
    cell.font = Font(
        color="0563C1",
        underline="single",
    )


def save_application(
    job,
    analysis,
    status,
    current_url="",
    external_url="",
    reason="",
    notes="",
    unanswered_question="",
):
    create_excel_file()

    workbook = load_workbook(EXCEL_FILE)
    sheet = workbook["Applications"]

    # Make sure header is current.
    for column, header in enumerate(HEADERS, start=1):
        sheet.cell(row=1, column=column, value=header)

    now = datetime.now()

    row = [
        now.strftime("%Y-%m-%d"),
        now.strftime("%H:%M:%S"),
        job.get("title", ""),
        job.get("company", ""),
        analysis.get("match_score", ""),
        analysis.get("decision", ""),
        status,
        job.get("url", ""),
        external_url,
        current_url,
        reason or analysis.get("reason", ""),
        notes,
        unanswered_question,
    ]

    sheet.append(row)

    current_row = sheet.max_row

    for column in range(1, len(HEADERS) + 1):
        sheet.cell(
            row=current_row,
            column=column,
        ).alignment = Alignment(
            vertical="top",
            wrap_text=True,
        )

    add_hyperlink(
        sheet.cell(row=current_row, column=8),
        job.get("url", ""),
    )

    add_hyperlink(
        sheet.cell(row=current_row, column=9),
        external_url,
    )

    add_hyperlink(
        sheet.cell(row=current_row, column=10),
        current_url,
    )

    workbook.save(EXCEL_FILE)

    print(
        f"Excel updated: "
        f"{job.get('title', '')} | "
        f"{job.get('company', '')} | "
        f"{status}"
    )


def save_skipped_job(
    job,
    analysis,
    reason="Job is not relevant",
):
    save_application(
        job=job,
        analysis=analysis,
        status="SKIP",
        reason=reason,
        notes="Job skipped by AI relevance system.",
    )


if __name__ == "__main__":
    prepare_real_tracker()
    print("Excel tracker is ready.")